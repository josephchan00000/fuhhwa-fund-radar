"""泡沫頁 VER 4「未來獲利」的資料抓取（docs/lab_risk_plan.md 第 11 節）。只在 CI 跑（本容器的 proxy 擋這些站）。

使用者 2026-10-18：股價反應的是未來獲利，要看未來獲利的成長，邊際成長放緩是警訊；那斯達克要有自己的基本面。
現在的等級用分析師預估，過去的驗證用替代資料（S&P 500 用專業預測者調查，真的預估）。

這一版只有探針 probe()：確認來源抓得到、格式長怎樣，再寫規格（同 U1 探針 → U2 規格的順序）。
"""
from __future__ import annotations

import io
import json
import re
import urllib.parse
import urllib.request

from ..http import _CTX
from .heatfetch import BROWSER, _get
from .market import log

SPF_PAGES = ("https://www.philadelphiafed.org/surveys-and-data/real-time-data-research/median-forecasts",
             "https://www.philadelphiafed.org/surveys-and-data/real-time-data-research/mean-forecasts",
             "https://www.philadelphiafed.org/surveys-and-data/cprof")
DBN_SEARCH = "https://api.db.nomics.world/v22/series/{p}?q={q}&observations=false&limit=80&format=json"
DBN_ONE = "https://api.db.nomics.world/v22/series/{p}?observations=1&format=json"
NASDAQ_HDR = {**BROWSER, "Accept": "application/json, text/plain, */*", "Origin": "https://www.nasdaq.com",
              "Referer": "https://www.nasdaq.com/"}


def _p(tag: str, msg: str) -> None:
    log(f"[fwd-probe] {tag} {msg}")


def _try(tag: str, url: str, *, timeout: int = 60, headers: dict | None = None, n: int = 400) -> bytes | None:
    try:
        b = _get(url, timeout=timeout, headers=headers)
        _p(tag, f"OK {len(b)} bytes：{b[:n]!r}")
        return b
    except Exception as e:                                    # noqa: BLE001 —— 探針：印出原因、往下試
        _p(tag, f"✗ {type(e).__name__} {str(e)[:200]}")
        return None


def _post(tag: str, url: str, form: dict, *, headers: dict | None = None, timeout: int = 60, n: int = 400) -> bytes | None:
    try:
        req = urllib.request.Request(url, data=urllib.parse.urlencode(form).encode(),
                                     headers={**(headers or BROWSER), "Content-Type": "application/x-www-form-urlencoded"})
        with urllib.request.urlopen(req, timeout=timeout, context=_CTX) as r:
            b = r.read()
        _p(tag, f"OK {len(b)} bytes：{b[:n]!r}")
        return b
    except Exception as e:                                    # noqa: BLE001
        _p(tag, f"✗ {type(e).__name__} {str(e)[:200]}")
        return None


def _xlsx_peek(tag: str, b: bytes, want: str = "CPROF") -> None:
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(b), read_only=True, data_only=True)
    except Exception as e:                                    # noqa: BLE001
        _p(tag, f"xlsx 讀不開 {type(e).__name__} {str(e)[:160]}")
        return
    names = wb.sheetnames
    _p(tag, f"工作表 {len(names)} 個：{names[:60]}")
    ws = next((wb[s] for s in names if s.strip().upper() == want), None) or wb.worksheets[0]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    _p(tag, f"工作表 {ws.title}：{len(rows)} 列")
    for r in rows[:4] + [["…"]] + rows[-4:]:
        _p(tag, f"  {r[:16]}")


def probe_spf() -> None:
    links: list[str] = []
    for page in SPF_PAGES:
        b = _try("SPF 頁", page, n=120)
        if not b:
            continue
        hrefs = re.findall(r'href="([^"]+\.xlsx[^"]*)"', b.decode("utf-8", "replace"), re.I)
        _p("SPF 頁", f"{page} 的 xlsx 連結 {len(hrefs)} 個：{[h[-90:] for h in hrefs[:30]]}")
        links += [urllib.parse.urljoin("https://www.philadelphiafed.org/", h.replace("&amp;", "&")) for h in hrefs]
    seen = set()
    for url in links:
        key = url.split("?")[0].lower()
        if key in seen or not re.search(r"(median|mean).?(level|growth)|cprof", key):
            continue
        seen.add(key)
        b = _try("SPF 檔", url, timeout=120, n=60)
        if b:
            _xlsx_peek("SPF 檔 " + key.rsplit("/", 1)[-1], b)


def probe_bea_tech() -> None:
    for ds in ("BEA/NIPA-T61600D", "BEA/NIPA-T61600C", "BEA/NIPA-T61600B"):
        for q in ("computer", "information", "electronic"):
            try:
                doc = json.loads(_get(DBN_SEARCH.format(p=ds, q=q), timeout=120).decode("utf-8"))
                docs = ((doc.get("series") or {}).get("docs")) or []
                _p("BEA", f"{ds} q={q}：{len(docs)} 條")
                for x in docs[:20]:
                    _p("BEA", f"   {x.get('series_code')}｜{str(x.get('series_name'))[:140]}")
            except Exception as e:                            # noqa: BLE001
                _p("BEA", f"{ds} q={q} ✗ {type(e).__name__} {str(e)[:160]}")


def probe_bea_series(codes: list[str]) -> None:
    for c in codes:
        try:
            doc = json.loads(_get(DBN_ONE.format(p=c), timeout=120).decode("utf-8"))
            d = ((doc.get("series") or {}).get("docs") or [{}])[0]
            per, val = d.get("period") or [], d.get("value") or []
            _p("BEA 序列", f"{c}：{len(per)} 期 {per[:2]}…{per[-2:]}，最後值 {val[-3:]}")
        except Exception as e:                                # noqa: BLE001
            _p("BEA 序列", f"{c} ✗ {type(e).__name__} {str(e)[:160]}")


def _wiki_tickers(tag: str, url: str) -> None:
    b = _try(tag, url, n=80)
    if not b:
        return
    html = b.decode("utf-8", "replace")
    m = re.search(r'<table[^>]*id="constituents".*?</table>', html, re.S)
    _p(tag, f"constituents 表 {'有' if m else '沒有'}；整頁表格 {len(re.findall(r'<table', html))} 個")
    seg = m.group(0) if m else html
    rows = re.findall(r"<tr[^>]*>(.*?)</tr>", seg, re.S)
    cells = [re.findall(r"<t[dh][^>]*>(.*?)</t[dh]>", r, re.S) for r in rows]
    clean = [[re.sub(r"<[^>]+>", "", c).strip()[:30] for c in r] for r in cells if r]
    _p(tag, f"列數 {len(clean)}；前 4 列 {clean[:4]}；最後 2 列 {clean[-2:]}")


def probe_constituents() -> None:
    _wiki_tickers("維基 S&P 500", "https://en.wikipedia.org/wiki/List_of_S%26P_500_companies")
    _wiki_tickers("維基 那斯達克 100", "https://en.wikipedia.org/wiki/Nasdaq-100")
    _wiki_tickers("維基 費半", "https://en.wikipedia.org/wiki/PHLX_Semiconductor_Sector")
    b = _try("SPY 持股", "https://www.ssga.com/us/en/intermediary/library-content/products/fund-data/etfs/us/"
             "holdings-daily-us-en-spy.xlsx", timeout=90, n=60)
    if b:
        _xlsx_peek("SPY 持股", b, want="HOLDINGS")
    _try("QQQ 持股", "https://www.invesco.com/us/financial-products/etfs/holdings/main/holdings/0"
         "?audienceType=Investor&action=download&ticker=QQQ", timeout=90, n=500)
    _try("那斯達克 100 API", "https://api.nasdaq.com/api/quote/list-type/nasdaq100", headers=NASDAQ_HDR, n=500)
    b = _try("那斯達克篩選器", "https://api.nasdaq.com/api/screener/stocks?tableonly=true&limit=6000&exchange=NASDAQ&download=true",
             headers=NASDAQ_HDR, timeout=120, n=300)
    if b:
        try:
            rows = (json.loads(b.decode("utf-8")).get("data") or {}).get("rows") or []
            _p("那斯達克篩選器", f"{len(rows)} 列；第一列 {rows[0] if rows else None}")
        except Exception as e:                                # noqa: BLE001
            _p("那斯達克篩選器", f"解析 ✗ {type(e).__name__} {str(e)[:120]}")
    _try("費半 指數頁", "https://indexes.nasdaq.com/Index/Weighting/SOX", n=300)
    _post("費半 權重 API", "https://indexes.nasdaqomx.com/Index/WeightingData",
          {"id": "SOX", "tradeDate": "2026-10-16T00:00:00.000", "timeOfDay": "SOD"}, n=600)
    _try("SOXX 持股", "https://www.ishares.com/us/products/239705/ishares-semiconductor-etf/1467271812596.ajax"
         "?fileType=csv&fileName=SOXX_holdings&dataType=fund", timeout=90, n=900)


def probe_yahoo_us() -> None:
    from . import fcst
    Y = fcst._Yahoo()
    if not Y.login():
        _p("Yahoo", "拿不到 crumb")
        return
    for sym in ("AAPL", "NVDA", "MSFT", "AVGO", "BRK-B"):
        st, rec = Y.one(sym)
        _p("Yahoo", f"{sym} {st}：{json.dumps(rec, ensure_ascii=False)[:700] if rec else None}")
    crumb = urllib.parse.quote(Y.crumb or "")
    st, b = Y.get(f"https://query2.finance.yahoo.com/v10/finance/quoteSummary/AAPL?modules=price,defaultKeyStatistics"
                  f"&crumb={crumb}")
    _p("Yahoo price", f"{st} {b[:900]!r}")
    st, b = Y.get(f"https://query2.finance.yahoo.com/v7/finance/quote?symbols=AAPL,NVDA,MSFT,AMD,INTC&crumb={crumb}")
    _p("Yahoo 批次報價", f"{st} {b[:1200]!r}")


def probe() -> None:
    for fn in (probe_spf, probe_bea_tech, probe_constituents, probe_yahoo_us):
        try:
            fn()
        except Exception as e:                                # noqa: BLE001
            _p(fn.__name__, f"✗ {type(e).__name__} {str(e)[:200]}")


def probe2() -> None:
    """第二輪：BEA 序列期間、SIC 舊分類的機械與通訊、SPF 年度預測從哪一季開始、ADR 與外幣公司的 Yahoo 欄位。"""
    probe_bea_series(["BEA/NIPA-T61600D/N501RC-Q", "BEA/NIPA-T61600D/N502RC-Q",
                      "BEA/NIPA-T61600C/A410RC-Q", "BEA/NIPA-T61600B/A410RBC-Q"])
    for ds in ("BEA/NIPA-T61600C", "BEA/NIPA-T61600B"):
        for q in ("machinery", "communication", "telephone", "services"):
            try:
                doc = json.loads(_get(DBN_SEARCH.format(p=ds, q=q), timeout=120).decode("utf-8"))
                docs = ((doc.get("series") or {}).get("docs")) or []
                _p("BEA", f"{ds} q={q}：{len(docs)} 條")
                for x in docs[:12]:
                    _p("BEA", f"   {x.get('series_code')}｜{str(x.get('series_name'))[:140]}")
            except Exception as e:                            # noqa: BLE001
                _p("BEA", f"{ds} q={q} ✗ {type(e).__name__} {str(e)[:160]}")
    try:
        import openpyxl
        page = _get(SPF_PAGES[2], timeout=60).decode("utf-8", "replace")
        href = next(h for h in re.findall(r'href="([^"]+\.xlsx[^"]*)"', page, re.I) if "median_cprof_level" in h.lower())
        b = _get(urllib.parse.urljoin("https://www.philadelphiafed.org/", href.replace("&amp;", "&")), timeout=120)
        rows = [list(r) for r in openpyxl.load_workbook(io.BytesIO(b), read_only=True, data_only=True).worksheets[0]
                .iter_rows(values_only=True)]
        hdr = rows[0]
        ia, ib = hdr.index("CPROFA"), hdr.index("CPROFB")
        first_a = next((r[:2] for r in rows[1:] if isinstance(r[ia], (int, float))), None)
        first_b = next((r[:2] for r in rows[1:] if isinstance(r[ib], (int, float))), None)
        gaps = [r[:2] for r in rows[1:] if not isinstance(r[3], (int, float))]
        _p("SPF", f"CPROFA 從 {first_a}、CPROFB 從 {first_b}；CPROF2 缺值的季 {gaps[:20]}")
        for r in rows[1:]:
            if r[0] in (2005.0, 2006.0) or (r[0] == 1981.0):
                _p("SPF", f"  {r}")
    except Exception as e:                                    # noqa: BLE001
        _p("SPF", f"✗ {type(e).__name__} {str(e)[:200]}")
    from . import fcst
    Y = fcst._Yahoo()
    if not Y.login():
        _p("Yahoo", "拿不到 crumb")
        return
    for sym in ("TSM", "ASML", "PDD", "ARM", "AZN"):
        st, rec = Y.one(sym)
        _p("Yahoo", f"{sym} {st}：{json.dumps(rec, ensure_ascii=False)[:600] if rec else None}")
    crumb = urllib.parse.quote(Y.crumb or "")
    st, b = Y.get(f"https://query2.finance.yahoo.com/v7/finance/quote?symbols=TSM,ASML,PDD,ARM,AZN,AAPL&crumb={crumb}")
    try:
        res = json.loads(b.decode("utf-8"))["quoteResponse"]["result"]
        keys = ("symbol", "currency", "financialCurrency", "marketCap", "sharesOutstanding", "regularMarketPrice",
                "epsTrailingTwelveMonths", "trailingPE", "epsCurrentYear", "priceEpsCurrentYear", "epsForward", "forwardPE")
        for q in res:
            _p("Yahoo 批次", str({k: q.get(k) for k in keys}))
    except Exception as e:                                    # noqa: BLE001
        _p("Yahoo 批次", f"{st} ✗ {type(e).__name__} {str(e)[:120]} {b[:300]!r}")
    big = ",".join(["AAPL", "MSFT", "NVDA", "GOOGL", "AMZN", "META", "AVGO", "TSLA", "BRK-B", "JPM"] * 15)
    st, b = Y.get(f"https://query2.finance.yahoo.com/v7/finance/quote?symbols={big}&crumb={crumb}")
    _p("Yahoo 批次 150", f"{st} {len(b)} bytes")
