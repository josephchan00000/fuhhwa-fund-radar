"""景氣溫度・美股版與台股泡沫面向的資料抓取（docs/lab_heat_us_plan.md 第 14、15 節；來源由 CI 探針四輪確認）。

只在 CI 跑（本容器的 proxy 擋這些站）。每個來源各自失敗、各自降級：抓不到就保留既有的檔，不把有值的檔洗成空。
存檔都走 heatdata.save()：data/lab/heat/{名稱}.json.gz，{"src", "d": {日期或 YYYY-MM: 值}}。
解析函式（parse_*）不碰網路，tests/test_heatfetch_us.py 用探針看到的格式測。

  us_px_{spx,ixic,ndx,sox}  Yahoo 指數日線收盤                    {日期: 收盤}           1927-12-30／1971／1985／1994 起
  us_vxn、us_vix3m、us_dxy  Yahoo ^VXN、^VIX3M、DX-Y.NYB 收盤      {日期: 收盤}           2001／2006／1971 起（VIX 用台股頁的 vix）
  us_rsp、us_spy            Yahoo 含息還原收盤（參考讀數）          {日期: 還原收盤}       2003／1993 起
  us_qqq、us_qqew           Yahoo 含息還原收盤（寬度與深度的大型股撐盤，breadthspec）  {日期: 還原收盤}  1999／2006 起
  us_h15                    聯準會 H.15（DBnomics）                 {日期: [3 個月國庫券, 3 個月固定期限, 2 年, 10 年]}  1954 起
  us_h41                    聯準會 H.4.1（DBnomics）                {週三: [總資產, 財政部存款, 逆回購（其他）]}（百萬美元）  2002-12 起
  us_m2                     聯準會 H.6 M2（季調，DBnomics）         {YYYY-MM: 十億美元}    1959 起
  us_z1                     聯準會 Z.1（DBnomics）                  {季底 YYYY-MM: [非金融企業股票市值, 非金融企業淨值, 家計證券融資]}（百萬美元）
  us_gdp                    BEA 名目 GDP（季調年率，DBnomics）      {季底 YYYY-MM: 百萬美元}
  us_cp                     BEA 企業稅後獲利（季、年率，DBnomics）  {季底 YYYY-MM: 十億美元}  泡沫頁 VER 2
  us_cftc_am、us_cftc_vix   CFTC 交易者分類：資產管理人 E-mini S&P、槓桿基金 VIX 期貨  {週二: [多單, 空單, 未平倉]}
  us_cftc_big、us_cftc_mini CFTC 傳統：非商業 S&P 大型、E-mini S&P {週二: [多單, 空單, 未平倉]}
  us_cftc_am_nq             CFTC 交易者分類：資產管理人 E-mini 那斯達克 100（heatspec_us VER 3 的 S8，2006-06 起）
  us_cftc_big_nq、us_cftc_mini_nq  CFTC 傳統：非商業 那斯達克 100 大型（1996 起、2015-06 下市）、E-mini（1999 起）（S9）
  us_lev_{代號}             Yahoo 3 倍槓桿 ETF 每天成交金額（收盤 × 成交量；分割前後不變）{日期: 美元}（S6）
  us_philly                 費城聯準製造業 目前一般活動             {YYYY-MM: 擴散指數}    1968-05 起
  us_ads                    費城聯準 ADS 每日景氣指數               {日期: 指數}           1960-03 起
  us_cli                    OECD 美國綜合領先指標（振幅調整）       {YYYY-MM: 指數}        1955 起
  us_claims                 勞工部 各州初領失業救濟金加總（未季調） {週六: 人數}           1986-02 起
  us_bls                    勞工統計局 失業率、CPI、核心 CPI        {YYYY-MM: [失業率, CPI, 核心 CPI]}  1947 起
  us_umich                  密西根大學消費者信心                    {YYYY-MM: 指數}        1952 起（1978 起每月）
  us_ebp                    聯準會 超額債券溢酬                      {YYYY-MM: [GZ 利差, 超額債券溢酬]}  1973 起
  us_shiller                Shiller 資料檔                           {YYYY-MM: [P 月平均, D, E, CPI, GS10]}  1871 起
  us_wsts                   WSTS 全球半導體銷售（千美元）            {YYYY-MM: 金額}        1986 起
  taiex90                   證交所 FMTQIK 加權指數收盤（1990～1998） {日期: 收盤}           1990-01-04 起（台股頁的 turnover 從 1999 起）
  twgdp                     主計總處 名目 GDP（季，百萬元）          {季底 YYYY-MM: 金額}   1981 起
"""
from __future__ import annotations

import csv
import io
import json
import re
import time
import urllib.parse
import urllib.request
from datetime import date, datetime, timedelta, timezone

from . import heatdata as HD
from .heatfetch import BROWSER, BLS, DBN, _get, _merge, _post_json, _replace, _twca_ctx, parse_dbnomics
from .market import log

# Yahoo 盤中的 bar：交易時段結束後 15 分鐘內都當成還沒收盤
CLOSE_MARGIN_S = 900
US_INDEX_SYM = {"spx": "^GSPC", "ixic": "^IXIC", "ndx": "^NDX", "sox": "^SOX"}


# ── Yahoo（只收已收盤的 bar；日期用交易所當地日期） ─────────────────────────────

def parse_yahoo(doc: dict, now: float, adj: bool = False) -> dict:
    """chart API 的回應 → {日期: 收盤（adj＝True 時含息還原收盤）}。
    日期＝時間戳加交易所時差；交易時段還沒結束（＋15 分鐘）的當天 bar 不收；週末 bar 不收。"""
    res = (((doc or {}).get("chart") or {}).get("result") or [None])[0]
    if not res:
        return {}
    meta = res.get("meta") or {}
    off = int(meta.get("gmtoffset") or -18000)
    reg = (meta.get("currentTradingPeriod") or {}).get("regular") or {}
    open_day = None
    if reg.get("start") and reg.get("end") and reg["start"] <= now < reg["end"] + CLOSE_MARGIN_S:
        open_day = datetime.fromtimestamp(int(reg["start"]) + off, timezone.utc).date()
    today = datetime.fromtimestamp(now + off, timezone.utc).date()
    ts = res.get("timestamp") or []
    if adj:
        vals = ((((res.get("indicators") or {}).get("adjclose") or [{}])[0]) or {}).get("adjclose") or []
    else:
        vals = ((((res.get("indicators") or {}).get("quote") or [{}])[0]) or {}).get("close") or []
    out = {}
    for t, v in zip(ts, vals):
        if v is None:
            continue
        d = datetime.fromtimestamp(int(t) + off, timezone.utc).date()
        if d.weekday() >= 5 or d > today or d == open_day:
            continue
        out[d.isoformat()] = round(float(v), 6 if adj else 4)
    return out


def yahoo_us(name: str, sym: str, start: str = "1900-01-01", adj: bool = False) -> int:
    p1 = int(datetime.fromisoformat(start).replace(tzinfo=timezone.utc).timestamp())
    p2 = int(time.time())
    url = (f"https://query1.finance.yahoo.com/v8/finance/chart/{urllib.parse.quote(sym)}"
           f"?period1={p1}&period2={p2}&interval=1d" + ("&events=div%2Csplit" if adj else ""))
    doc = json.loads(_get(url, timeout=60).decode("utf-8"))
    new = parse_yahoo(doc, time.time(), adj=adj)
    if adj:                                               # 含息還原會改寫整段歷史：整段換
        return _replace(name, "yahoo", new)
    return _merge(name, "yahoo", new)


LEV_START = "2008-11-01"                                   # 最早的一檔（SPXL、SPXS）2008-11 上市


def parse_yahoo_dv(doc: dict, now: float) -> dict:
    """chart API → {日期: 成交金額（收盤 × 成交量，美元）}；日期與還沒收盤的判斷同 parse_yahoo。
    Yahoo 的收盤與成交量都會依之後的分割調整，相乘後和當時的成交金額一樣，所以歷史不會因為分割改寫。"""
    res = (((doc or {}).get("chart") or {}).get("result") or [None])[0]
    if not res:
        return {}
    meta = res.get("meta") or {}
    off = int(meta.get("gmtoffset") or -18000)
    reg = (meta.get("currentTradingPeriod") or {}).get("regular") or {}
    open_day = None
    if reg.get("start") and reg.get("end") and reg["start"] <= now < reg["end"] + CLOSE_MARGIN_S:
        open_day = datetime.fromtimestamp(int(reg["start"]) + off, timezone.utc).date()
    today = datetime.fromtimestamp(now + off, timezone.utc).date()
    q = ((((res.get("indicators") or {}).get("quote") or [{}])[0]) or {})
    out = {}
    for t, c, v in zip(res.get("timestamp") or [], q.get("close") or [], q.get("volume") or []):
        if c is None or v is None:
            continue
        d = datetime.fromtimestamp(int(t) + off, timezone.utc).date()
        if d.weekday() >= 5 or d > today or d == open_day:
            continue
        out[d.isoformat()] = round(float(c) * float(v), 0)
    return out


def lev_names() -> list[str]:
    from . import heatspec_us as HU
    return sorted({s for pair in HU.LEV.values() for s in pair})


def yahoo_dv(name: str, sym: str, start: str = LEV_START) -> int:
    """槓桿 ETF 的每天成交金額（heatspec_us VER 3 的 S6）。Yahoo 偶爾回 429：重試兩次、每次多等一點。"""
    p1 = int(datetime.fromisoformat(start).replace(tzinfo=timezone.utc).timestamp())
    url = (f"https://query1.finance.yahoo.com/v8/finance/chart/{urllib.parse.quote(sym)}"
           f"?period1={p1}&period2={int(time.time())}&interval=1d")
    err = None
    for k in range(3):
        try:
            doc = json.loads(_get(url, timeout=60).decode("utf-8"))
            return _merge(name, "yahoo", parse_yahoo_dv(doc, time.time()))
        except Exception as e:                                # noqa: BLE001 —— 429 等一下再試
            err = e
            time.sleep(5 * (k + 1))
    raise err


# ── DBnomics（聯準會 H.15、H.4.1、H.6、Z.1、BEA） ───────────────────────────────

def dbn_series(path: str) -> dict:
    return parse_dbnomics(json.loads(_get(DBN.format(p=path), timeout=120).decode("utf-8")))


def month_key(p: str) -> str | None:
    """DBnomics 的期別 → YYYY-MM：'1959-01-31'、'1959-01'、'1947-Q1'（季底月份）。"""
    m = re.match(r"^(\d{4})-(\d{2})", p)
    if m:
        return f"{m.group(1)}-{m.group(2)}"
    m = re.match(r"^(\d{4})-Q([1-4])$", p)
    if m:
        return f"{m.group(1)}-{int(m.group(2)) * 3:02d}"
    return None


def combine(series: list[dict], key=lambda k: k) -> dict:
    """好幾條序列 → {鍵: [各序列的值或 None]}（至少一個有值才留）。"""
    out: dict = {}
    for j, s in enumerate(series):
        for p, v in s.items():
            k = key(p)
            if k is None:
                continue
            out.setdefault(k, [None] * len(series))[j] = v
    return out


H15 = ("FED/H15/RIFSGFSM03_N.B", "FED/H15/RIFLGFCM03_N.B", "FED/H15/RIFLGFCY02_N.B", "FED/H15/RIFLGFCY10_N.B")
H41 = ("FED/H41/RESPPMA_N.WW", "FED/H41/RESPPLLDT_N.WW", "FED/H41/RESPPLLRD_N.WW")
Z1 = ("FED/Z1/LM103164105.Q", "FED/Z1/FL102090005.Q", "FED/Z1/FL153067005.Q")
M2 = "FED/H6_H6_M2/M2.M"
GDP = "BEA/NIPA-T10105/A191RC-Q"


def h15() -> int:
    day = lambda p: p if re.match(r"^\d{4}-\d{2}-\d{2}$", p) else None  # noqa: E731
    return _merge("us_h15", "fed_h15_dbnomics", combine([dbn_series(p) for p in H15], day))


def h41() -> int:
    day = lambda p: p if re.match(r"^\d{4}-\d{2}-\d{2}$", p) else None  # noqa: E731
    return _merge("us_h41", "fed_h41_dbnomics", combine([dbn_series(p) for p in H41], day))


def m2() -> int:
    return _merge("us_m2", "fed_h6_dbnomics", {month_key(p): v for p, v in dbn_series(M2).items() if month_key(p)})


# 美國企業稅後獲利（BEA NIPA，季資料、年率、十億美元；泡沫頁 VER 2 那斯達克兩檔的基本面近似，docs/lab_risk_plan.md 第 9 節）
# 第一順位＝不含存貨評價與資本消耗調整（A055RC，最接近公司財報的盈餘）；DBnomics 沒有才用含調整的（A551RC）
# 2026-10-18 CI 探針：表 1.12 第 45 行 A055RC-Q 有季資料（1947 起）；表 6.19D 的 A055RC 只有年資料
CP_CANDS = ("BEA/NIPA-T11200/A055RC-Q", "BEA/NIPA-T11200/A551RC-Q")
DBN_SEARCH = "https://api.db.nomics.world/v22/series/{p}?q={q}&observations=false&limit=60&format=json"


def cp() -> int:
    """依 CP_CANDS 順序取第一條有資料的；用了哪一條記在來源標籤（us_cp 的 src）。"""
    errs = []
    for sid in CP_CANDS:
        try:
            d = {month_key(p): v for p, v in dbn_series(sid).items() if month_key(p) and v is not None}
        except Exception as e:                                # noqa: BLE001 —— 換下一條
            errs.append(f"{sid}：{type(e).__name__} {str(e)[:80]}")
            continue
        if len(d) >= 200:
            log(f"[heat-us] 美國企業獲利用 {sid}（{min(d)}～{max(d)}，{len(d)} 季）")
            return _merge("us_cp", "bea_dbnomics:" + sid, d)
        errs.append(f"{sid}：只有 {len(d)} 筆")
    raise RuntimeError("美國企業獲利都抓不到：" + "；".join(errs))


def probe_cp() -> None:
    """探針：列出 BEA 幾張表裡名稱有 profits 的序列（代碼＋名稱），決定 CP_CANDS 用哪一條。"""
    for ds in ("BEA/NIPA-T11200", "BEA/NIPA-T61900D", "BEA/NIPA-T61600D", "BEA/NIPA-T11400"):
        try:
            doc = json.loads(_get(DBN_SEARCH.format(p=ds, q=urllib.parse.quote("profits after tax")), timeout=120).decode("utf-8"))
            docs = ((doc.get("series") or {}).get("docs")) or []
            log(f"[cp-probe] {ds}：{len(docs)} 條")
            for x in docs[:40]:
                log(f"[cp-probe]   {x.get('series_code')}｜{str(x.get('series_name'))[:120]}")
        except Exception as e:                                # noqa: BLE001
            log(f"[cp-probe] {ds} ✗ {type(e).__name__} {str(e)[:160]}")


def z1() -> int:
    return _merge("us_z1", "fed_z1_dbnomics", combine([dbn_series(p) for p in Z1], month_key))


def gdp() -> int:
    return _merge("us_gdp", "bea_dbnomics", {month_key(p): v for p, v in dbn_series(GDP).items() if month_key(p)})


# ── CFTC（Socrata；每週二部位） ───────────────────────────────────────────────

CFTC = "https://publicreporting.cftc.gov/resource/{ds}.json"
CFTC_JOBS = {   # 檔名: (資料集, 契約代碼, 多單欄, 空單欄)
    "us_cftc_am": ("gpe5-46if", "13874A", "asset_mgr_positions_long", "asset_mgr_positions_short"),
    "us_cftc_vix": ("gpe5-46if", "1170E1", "lev_money_positions_long", "lev_money_positions_short"),
    "us_cftc_big": ("6dca-aqww", "138741", "noncomm_positions_long_all", "noncomm_positions_short_all"),
    "us_cftc_mini": ("6dca-aqww", "13874A", "noncomm_positions_long_all", "noncomm_positions_short_all"),
    # heatspec_us VER 3：那斯達克 100（E-mini 209742、大型契約 209741；docs/lab_us_chip_plan.md 第 7 節）
    "us_cftc_am_nq": ("gpe5-46if", "209742", "asset_mgr_positions_long", "asset_mgr_positions_short"),
    "us_cftc_big_nq": ("6dca-aqww", "209741", "noncomm_positions_long_all", "noncomm_positions_short_all"),
    "us_cftc_mini_nq": ("6dca-aqww", "209742", "noncomm_positions_long_all", "noncomm_positions_short_all"),
}


def _num(x):
    if x is None:
        return None
    try:
        s = str(x).replace(",", "").strip()
        return float(s) if s not in ("", ".", "NaN", "nan", "NA") else None
    except ValueError:
        return None


def parse_cftc(rows: list, lcol: str, scol: str) -> dict:
    out = {}
    for r in rows or []:
        d = str(r.get("report_date_as_yyyy_mm_dd") or "")[:10]
        lo, sh, oi = _num(r.get(lcol)), _num(r.get(scol)), _num(r.get("open_interest_all"))
        if re.match(r"^\d{4}-\d{2}-\d{2}$", d) and lo is not None and sh is not None and oi:
            out[d] = [lo, sh, oi]
    return out


def cftc(name: str, since: str = "1986-01-01") -> int:
    ds, code, lcol, scol = CFTC_JOBS[name]
    where = urllib.parse.quote(f"cftc_contract_market_code='{code}' AND report_date_as_yyyy_mm_dd >= '{since}T00:00:00'")
    sel = urllib.parse.quote(f"report_date_as_yyyy_mm_dd,open_interest_all,{lcol},{scol}")
    url = (CFTC.format(ds=ds) + f"?$select={sel}&$where={where}&$order=report_date_as_yyyy_mm_dd&$limit=50000")
    return _merge(name, "cftc", parse_cftc(json.loads(_get(url, timeout=120).decode("utf-8")), lcol, scol))


# ── 費城聯準、OECD、勞工部、密大、聯準會 EBP ─────────────────────────────────────

PHILLY = ("https://www.philadelphiafed.org/-/media/FRBP/Assets/Surveys-And-Data/MBOS/Historical-Data/"
          "Diffusion-Indexes/bos_dif.csv")
ADS_PAGE = "https://www.philadelphiafed.org/surveys-and-data/real-time-data-research/ads"
OECD_CLI = ("https://sdmx.oecd.org/public/rest/data/OECD.SDD.STES,DSD_STES@DF_CLI,4.1/USA.M.LI...AA...H"
            "?startPeriod=1950-01&format=csvfilewithlabels")
AR539 = "https://oui.doleta.gov/unemploy/csv/ar539.csv"
UMICH = "https://www.sca.isr.umich.edu/files/tbmics.csv"
EBP = "https://www.federalreserve.gov/econres/notes/feds-notes/ebp_csv.csv"
MON = {m: i for i, m in enumerate(("jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"), 1)}


def _rows(txt: str) -> list[list[str]]:
    return [r for r in csv.reader(io.StringIO(txt.lstrip("﻿"))) if any(c.strip() for c in r)]


def parse_philly(txt: str, col: str = "GAC") -> dict:
    """bos_dif.csv：DATE＝'May-68'（兩位數年：60 以上＝19xx）；col 欄＝目前一般活動。"""
    rows = _rows(txt)
    if not rows:
        return {}
    j = [h.strip().upper() for h in rows[0]].index(col)
    out = {}
    for r in rows[1:]:
        m = re.match(r"^([A-Za-z]{3})-(\d{2})$", r[0].strip())
        v = _num(r[j]) if j < len(r) else None
        if m and v is not None and m.group(1).lower() in MON:
            yy = int(m.group(2))
            out[f"{1900 + yy if yy >= 60 else 2000 + yy:04d}-{MON[m.group(1).lower()]:02d}"] = v
    return out


def parse_ads(rows: list) -> dict:
    """ADS 試算表的列（openpyxl values）：['Date', 'ADS_Index', ...]、['1960:03:01', -0.58, 0]。"""
    out = {}
    for r in rows:
        if not r:
            continue
        d, v = r[0], r[1] if len(r) > 1 else None
        if isinstance(d, datetime):
            k = d.date().isoformat()
        else:
            m = re.match(r"^(\d{4})[:\-/](\d{2})[:\-/](\d{2})$", str(d or "").strip())
            k = f"{m.group(1)}-{m.group(2)}-{m.group(3)}" if m else None
        if k and isinstance(v, (int, float)) and not isinstance(v, bool):
            out[k] = round(float(v), 6)
    return out


def ads() -> int:
    import openpyxl
    html = _get(ADS_PAGE, timeout=90).decode("utf-8", "replace")
    href = next((h for h in re.findall(r'href="([^"]+)"', html, re.I) if re.search(r"ADS_Index_Most_Current_Vintage\.xlsx", h, re.I)), None)
    if not href:
        raise RuntimeError("找不到 ADS 檔案連結")
    url = urllib.parse.urljoin("https://www.philadelphiafed.org/", href.replace("&amp;", "&"))
    wb = openpyxl.load_workbook(io.BytesIO(_get(url, timeout=120)), read_only=True, data_only=True)
    rows = [list(r)[:3] for r in wb.worksheets[0].iter_rows(values_only=True)]
    return _merge("us_ads", "philly_ads", parse_ads(rows))


def parse_oecd(txt: str) -> dict:
    out = {}
    for r in csv.DictReader(io.StringIO(txt.lstrip("﻿"))):
        k, v = str(r.get("TIME_PERIOD") or ""), _num(r.get("OBS_VALUE"))
        if re.match(r"^\d{4}-\d{2}$", k) and v is not None:
            out[k] = v
    return out


def parse_ar539(txt: str) -> dict:
    """各州週資料：st, rptdate（週六）, c1, c2, c3（初領，未季調）… → {週六: 全部州的初領合計}。"""
    rows = _rows(txt)
    if not rows:
        return {}
    hdr = [h.strip().lower() for h in rows[0]]
    jd, jv = hdr.index("rptdate"), hdr.index("c3")
    out: dict = {}
    for r in rows[1:]:
        if len(r) <= max(jd, jv):
            continue
        d, v = r[jd].strip(), _num(r[jv])
        if re.match(r"^\d{4}-\d{2}-\d{2}$", d) and v is not None:
            out[d] = out.get(d, 0.0) + v
    return out


def parse_umich(txt: str) -> dict:
    """tbmics.csv：Month（英文月名）, YYYY, ICS_ALL。"""
    rows = _rows(txt)
    out = {}
    for r in rows[1:]:
        if len(r) < 3:
            continue
        m, y, v = r[0].strip().lower()[:3], r[1].strip(), _num(r[2])
        if m in MON and re.match(r"^\d{4}$", y) and v is not None:
            out[f"{y}-{MON[m]:02d}"] = v
    return out


def parse_ebp(txt: str) -> dict:
    rows = _rows(txt)
    if not rows:
        return {}
    hdr = [h.strip().lower() for h in rows[0]]
    jd, jg, je = hdr.index("date"), hdr.index("gz_spread"), hdr.index("ebp")
    out = {}
    for r in rows[1:]:
        m = re.match(r"^(\d{4})-(\d{2})", r[jd].strip()) if len(r) > jd else None
        if m:
            g, e = _num(r[jg]), _num(r[je])
            if e is not None:
                out[f"{m.group(1)}-{m.group(2)}"] = [None if g is None else round(g, 6), round(e, 6)]
    return out


def philly() -> int:
    return _merge("us_philly", "philly_mbos", parse_philly(_get(PHILLY, timeout=90).decode("utf-8", "replace")))


def oecd_cli() -> int:
    return _merge("us_cli", "oecd", parse_oecd(_get(OECD_CLI, timeout=120).decode("utf-8", "replace")))


def claims() -> int:
    return _merge("us_claims", "dol_ar539", parse_ar539(_get(AR539, timeout=180).decode("utf-8", "replace")))


def umich() -> int:
    return _merge("us_umich", "umich", parse_umich(_get(UMICH, timeout=90).decode("utf-8", "replace")))


def ebp() -> int:
    return _merge("us_ebp", "fed_ebp", parse_ebp(_get(EBP, timeout=90).decode("utf-8", "replace")))


# ── 勞工統計局（失業率、CPI、核心 CPI） ─────────────────────────────────────────

BLS_SERIES = ("LNS14000000", "CUUR0000SA0", "CUUR0000SA0L1E")


def parse_bls(doc: dict) -> dict:
    out: dict = {}
    for s in ((doc.get("Results") or {}).get("series")) or []:
        sid = s.get("seriesID")
        if sid not in BLS_SERIES:
            continue
        j = BLS_SERIES.index(sid)
        for r in s.get("data") or []:
            p, v = str(r.get("period") or ""), _num(r.get("value"))
            if re.match(r"^M(0[1-9]|1[0-2])$", p) and v is not None:
                out.setdefault(f"{r.get('year')}-{p[1:]}", [None, None, None])[j] = v
    return out


def bls(y0: int, y1: int | None = None) -> int:
    """API v1（不用金鑰）：POST 一次最多 10 年；三條序列一起抓。"""
    y1 = y1 or date.today().year
    out: dict = {}
    a = y0
    while a <= y1:
        z = min(a + 9, y1)
        doc = _post_json(BLS, {"seriesid": list(BLS_SERIES), "startyear": str(a), "endyear": str(z)})
        if doc.get("status") != "REQUEST_SUCCEEDED":
            raise RuntimeError(f"BLS {a}–{z} {doc.get('status')} {str(doc.get('message'))[:120]}")
        for k, v in parse_bls(doc).items():
            row = out.setdefault(k, [None, None, None])
            for j, x in enumerate(v):
                if x is not None:
                    row[j] = x
        a = z + 1
        time.sleep(1)
    old = HD.load("us_bls")                               # 有些年份這次沒抓（每天只補最近一年）→ 保留既有欄位
    for k, v in out.items():
        prev = old.get(k) or [None, None, None]
        out[k] = [x if x is not None else p for x, p in zip(v, prev)]
    return _merge("us_bls", "bls", out)


# ── Shiller、WSTS ────────────────────────────────────────────────────────────

SHILLER = "https://shillerdata.com/"
WSTS_PAGE = "https://www.wsts.org/67/Historical-Billings-Report"


def shiller_month(x) -> str | None:
    """Shiller 的日期欄：1871.01＝1 月、2025.1＝10 月、2025.12＝12 月。"""
    if not isinstance(x, (int, float)) or isinstance(x, bool):
        return None
    y = int(x)
    m = int(round((x - y) * 100))
    return f"{y:04d}-{m:02d}" if 1 <= m <= 12 and 1800 < y < 2200 else None


def parse_shiller(rows: list) -> dict:
    """Data 工作表的列：Date, P, D, E, CPI, Date Fraction, GS10, … → {YYYY-MM: [P, D, E, CPI, GS10]}（空白＝None）。"""
    out = {}
    for r in rows:
        if not r:
            continue
        k = shiller_month(r[0])
        if not k:
            continue
        vals = [r[j] if j < len(r) else None for j in (1, 2, 3, 4, 6)]
        vals = [float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None for v in vals]
        if vals[0] is not None:
            out[k] = vals
    return out


def shiller() -> int:
    import xlrd
    html = _get(SHILLER, timeout=90).decode("utf-8", "replace")
    url = re.findall(r'href="([^"]+\.xls[^"]*)"', html, re.I)[0]
    url = ("https:" + url) if url.startswith("//") else urllib.parse.urljoin(SHILLER, url)
    wb = xlrd.open_workbook(file_contents=_get(url, timeout=120))
    sh = wb.sheet_by_name("Data")
    new = parse_shiller([sh.row_values(i) for i in range(sh.nrows)])
    if len(new) < 1500:
        raise RuntimeError(f"Shiller 只解析出 {len(new)} 個月")
    return _replace("us_shiller", "shiller", new)        # 最近幾個月的 D、E 會補上：整段換


def parse_wsts(rows: list) -> dict:
    """Monthly Data 工作表：一列年份（int）接各地區列；Worldwide 列的第 2～13 欄＝1～12 月（千美元）。"""
    out = {}
    year = None
    for r in rows:
        if not r:
            continue
        a = r[0]
        if isinstance(a, (int, float)) and not isinstance(a, bool) and 1980 < a < 2200 and float(a) == int(a):
            year = int(a)
            continue
        if year and isinstance(a, str) and a.strip().lower() == "worldwide":
            for m in range(1, 13):
                v = r[m] if m < len(r) else None
                if isinstance(v, (int, float)) and not isinstance(v, bool) and v > 0:
                    out[f"{year:04d}-{m:02d}"] = float(v)
    return out


def wsts() -> int:
    import openpyxl
    html = _get(WSTS_PAGE, timeout=90).decode("utf-8", "replace")
    href = re.findall(r'href="([^"]+\.xlsx[^"]*)"', html, re.I)[0]
    url = href if href.startswith("http") else urllib.parse.urljoin("https://www.wsts.org/", href)
    wb = openpyxl.load_workbook(io.BytesIO(_get(url, timeout=120)), read_only=True, data_only=True)
    ws = wb["Monthly Data"] if "Monthly Data" in wb.sheetnames else wb.worksheets[0]
    return _merge("us_wsts", "wsts", parse_wsts([list(r)[:14] for r in ws.iter_rows(values_only=True)]))


# ── 台股：1990～1998 加權指數、季 GDP ─────────────────────────────────────────────

DGBAS_GDP = "https://ws.dgbas.gov.tw/001/Upload/461/relfile/11525/230912/na8103a1q.xml"
GDP_ITEM = "國內生產毛額_當期價格(新臺幣百萬元)、GDP"


def parse_fmtqik(doc: dict) -> dict:
    out = {}
    for r in (doc or {}).get("data") or []:
        mm = re.match(r"^\s*(\d{2,3})/(\d{2})/(\d{2})", str(r[0]))
        if not mm or len(r) < 5:
            continue
        v = _num(r[4])
        if v:
            out[f"{int(mm.group(1)) + 1911:04d}-{mm.group(2)}-{mm.group(3)}"] = v
    return out


def taiex90(a: date = date(1990, 1, 1), z: date = date(1998, 12, 31), budget_s: float = 1200) -> int:
    """FMTQIK 每月一請求（1990-01-04 起；台股頁的 turnover 從 1999 起，這段另存 taiex90，不改台股頁的交易日曆）。"""
    have = {k[:7] for k in HD.load("taiex90")}
    out: dict = {}
    t0 = time.time()
    y, m = a.year, a.month
    while (y, m) <= (z.year, z.month):
        if time.time() - t0 > budget_s:
            log("[heat-us] taiex90 時間到，下一輪續補")
            break
        ym = f"{y:04d}-{m:02d}"
        if ym not in have:
            url = f"https://www.twse.com.tw/rwd/zh/afterTrading/FMTQIK?date={y}{m:02d}01&response=json"
            try:
                out.update(parse_fmtqik(json.loads(_get(url, timeout=40).decode("utf-8"))))
            except Exception as ex:                       # noqa: BLE001
                log(f"[heat-us] FMTQIK {ym} ✗ {str(ex)[:80]}")
            time.sleep(3.5)
        m += 1
        if m == 13:
            y, m = y + 1, 1
    return _merge("taiex90", "twse", out)


def parse_dgbas_gdp(txt: str) -> dict:
    out = {}
    for blk in re.findall(r"<Obs>(.*?)</Obs>", txt, re.S):
        it = re.search(r"<Item>(.*?)</Item>", blk)
        if not it or it.group(1).strip() != GDP_ITEM or "<TYPE>原始值" not in blk:
            continue
        tp = re.search(r"<TIME_PERIOD>(\d{4})Q([1-4])</TIME_PERIOD>", blk)
        v = re.search(r"<Item_VALUE>(.*?)</Item_VALUE>", blk)
        x = _num(v.group(1)) if v else None
        if tp and x is not None:
            out[f"{tp.group(1)}-{int(tp.group(2)) * 3:02d}"] = x
    return out


def parse_dgbas_real_yoy(txt: str) -> dict:
    """同一支 XML 的實質 GDP 年增率（經濟成長率；行事曆補值用，arm/calsrc.py）→ {季末月: 年增率}。參考年改了照樣認得。"""
    out = {}
    for blk in re.findall(r"<Obs>(.*?)</Obs>", txt, re.S):
        it = re.search(r"<Item>(.*?)</Item>", blk)
        if not it or not re.match(r"國內生產毛額_連鎖實質值.*、GDP$", it.group(1).strip()) or "<TYPE>年增率" not in blk:
            continue
        tp = re.search(r"<TIME_PERIOD>(\d{4})Q([1-4])</TIME_PERIOD>", blk)
        v = re.search(r"<Item_VALUE>(.*?)</Item_VALUE>", blk)
        x = _num(v.group(1)) if v else None
        if tp and x is not None:
            out[f"{tp.group(1)}-{int(tp.group(2)) * 3:02d}"] = x
    return out


def twgdp(timeout: int = 180) -> int:
    req = urllib.request.Request(DGBAS_GDP, headers=BROWSER)
    with urllib.request.urlopen(req, timeout=timeout, context=_twca_ctx()) as r:
        txt = r.read().decode("utf-8", "replace")
    return _merge("twgdp", "dgbas", parse_dgbas_gdp(txt)) + _merge("twgdp_yoy", "dgbas", parse_dgbas_real_yoy(txt))


def latest_us_close_day(now: datetime | None = None) -> str:
    """最近一個應該已經收盤的美股交易日（美東 16:15 以後算當天；週末往前推）。只用來決定要不要再抓 Yahoo。"""
    now = now or datetime.now(timezone.utc)
    et = now - timedelta(hours=5)                         # 用 EST 估（夏令時間會晚一小時才換日，偏保守）
    d = et.date() if et.hour * 60 + et.minute >= 16 * 60 + 15 else et.date() - timedelta(days=1)
    while d.weekday() >= 5:
        d -= timedelta(days=1)
    return d.isoformat()
